#Import the required libraries
import boto3
import os
import math
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from botocore.exceptions import ClientError
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication


#ADDITIONAL FUNCTIONS ----------------------------------------------------------------------------------------


#function to save DataFrame to S3
def save_df_to_s3(etl, trace, partial, incode, glos, s3_client, s3_bucket, s3_key):
    excel_buffer = BytesIO()

    wb = Workbook()

    ws = wb.active
    ws.title = "Combined Log"
    ws.append(['ETL LOG'])
    for row in dataframe_to_rows(etl, index=False, header=True):
        ws.append(row)
    ws.append([])
    ws.append(['TRANSACTION LOG'])
    for row in dataframe_to_rows(trace, index=False, header=True):
        ws.append(row)

    ws1 = wb.create_sheet(title="Daily Partial Trans")
    for row in dataframe_to_rows(partial, index=False, header=True):
        ws1.append(row)

    ws2 = wb.create_sheet(title="Internal Code & Glosary")
    ws2.append(['INTERNAL CODE'])
    for row in dataframe_to_rows(incode, index=False, header=True):
        ws2.append(row)
    ws2.append([])
    ws2.append(['GLOSARY'])
    for row in dataframe_to_rows(glos, index=False, header=True):
        ws2.append(row)
    
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    s3_client.put_object(
        Bucket=s3_bucket,
        Key=s3_key,
        Body=excel_buffer.getvalue()
    )

#function to download report from s3
def download_file_from_s3(s3_client, s3_bucket, s3_key):
    try:
        s3_response = s3_client.get_object(Bucket=s3_bucket, Key=s3_key)
        file_content = s3_response['Body'].read()
        return file_content
    except ClientError as e:
        print(f"Failed to download file from S3: {e.response['Error']['Message']}")
        raise e

#function to create internal code
def internal_code():
    internal_code = [
        ['TRANS_MB51DS', 'M-MB51DS01', 'Medium', 'gain or loss in MB51DS'],
        ['TRANS_MB51DS', 'M-MB51DS02', 'Medium', 'noise MB51DS'],
        ['TRANS_MB51DS', 'M-MB51DS03', 'Medium', 'MB51DS that will be process tomorrow'],
        ['TRANS_MB51DS', 'L-MB51DS01', 'Low', 'block-unblock process MB51DS'],
        ['TRANS_MB51US', 'M-MB51US01', 'Medium', 'gain or loss in MB51US'],
        ['TRANS_MB51US', 'M-MB51US02', 'Medium', 'noise MB51US'],
        ['TRANS_MB51US', 'M-MB51US03', 'Medium', 'MB51US that will be process tomorrow'],
        ['TRANS_MB51US', 'L-MB51US01', 'Low', 'record with unused material'],
        ['TRANS_MB51US', 'L-MB51US02', 'Low', 'plant estate with material FRESH FRUIT BUNCH (Not use on tracebility because use ZMB)'],
        ['TRANS_MB51US', 'L-MB51US03', 'Low', 'plant estate with ordertype not null'],
        ['TRANS_MB51US', 'L-MB51US04', 'Low', 'plant estate with unused movementtype'],
        ['TRANS_MB51US', 'L-MB51US05', 'Low', 'plant mill ZMB with material FRESH FRUIT BUNCH (Not use on tracebility because use ZMB)'],
        ['TRANS_MB51US', 'L-MB51US06', 'Low', 'plant mill ZMB with ordertype not null'],
        ['TRANS_MB51US', 'L-MB51US07', 'Low', 'plant mill ZMB unused'],
        ['TRANS_OPENLINK', 'H-OPNK01', 'High', 'Valid Trans ID SODOGI Connected to valid DSI Interco Sales does not exist in events_bc (affected by invalid logic)'],
        ['TRANS_OPENLINK', 'H-OPNK02', 'High', 'Valid Trans ID DSI Interco Sales does not exist in events_bc (affected by invalid logic)'],
        ['TRANS_OPENLINK', 'M-OPNK01', 'Medium', 'DSI Interco Sales with sap_so_po_number does not exist in precedingdocument trans_mb51_ds'],
        ['TRANS_OPENLINK', 'M-OPNK02', 'Medium', 'DSI Interco Sales with sap_so_po_number does not exist in sales_order_no trans_sodogi'],
        ['TRANS_OPENLINK', 'M-OPNK03', 'Medium', 'Sales Order DSI Interco Sales has invalid status in trans_sodogi'],
        ['TRANS_OPENLINK', 'M-OPNK04', 'Medium', 'DSI Interco Sales with sap_so_po_number = 0'],
        ['TRANS_OPENLINK', 'M-OPNK05', 'Medium', '3rd Party Sales with backstop_id does not exist in backstop_id DSI Interco Sales trans_openlink'],
        ['TRANS_OPENLINK', 'M-OPNK06', 'Medium', 'DSI Interco Sales connected to 3rd Party Sales has invalid status'],
        ['TRANS_OPENLINK', 'M-OPNK07', 'Medium', 'DSI Interco Sales connected to 3rd Party Sales has sales order with invalid status in trans_sodogi'],
        ['TRANS_OPENLINK', 'M-OPNK08', 'Medium', 'without bl_date'],
        ['TRANS_OPENLINK', 'M-OPNK09', 'Medium', '3rd Party Sales with backstop_id already processed to events_bc. this data will be processed once versioning is implemented'],
        ['TRANS_OPENLINK', 'M-OPNK10', 'Medium', 'DSI Interco Sales with sap_so_po_number already processed to events_bc. this data will be processed once versioning is implemented'],
        ['TRANS_OPENLINK', 'L-OPNK01', 'Low', 'delivery_status is not Approved'],
        ['TRANS_OPENLINK', 'L-OPNK02', 'Low', 'contract_type `GAR Interco Purchase`'],
        ['TRANS_OPENLINK', 'L-OPNK03', 'Low', 'backstop_id = 0'],
        ['TRANS_OPENLINK', 'L-OPNK04', 'Low', 'status `invalid` (Openlink updated by new data)'],
        ['TRANS_SODOGIDS', 'M-SODOGI01', 'Medium', 'do_qty = 0'],
        ['TRANS_SODOGIDS', 'M-SODOGI02', 'Medium', 'without do_number'],
        ['TRANS_SODOGIDS', 'M-SODOGI03', 'Medium', 'without batch_no'],
        ['TRANS_SODOGIDS', 'M-SODOGI04', 'Medium', 'sales_order_no = 0'],
        ['TRANS_SODOGIDS', 'M-SODOGI05', 'Medium', 'sales_order_no does not exist in precedingdocument trans_mb51_ds'],
        ['TRANS_SODOGIDS', 'M-SODOGI06', 'Medium', 'sodogi data with the combination of sales_order_no and batch has been loaded into events_bc on previous days'],
        ['TRANS_SODOGIDS', 'M-SODOGI07', 'Medium', 'sodogi data with the same sales_order_no has been processed in events_bc. waiting for this batch to be processed'],
        ['TRANS_SODOGIDS', 'M-SODOGI08', 'Medium', 'without actual_gi_date'],
        ['TRANS_SODOGIDS', 'M-SODOGI09', 'Medium', 'SODOGI that will be process tomorrow or in the next few days'],
        ['TRANS_SODOGIDS', 'M-SODOGI10', 'Medium', 'MB51DS mvt.601 has already been connected to MB51DS mvt.101 in the previous Traceability'],
        ['TRANS_SODOGIDS', 'L-SODOGI01', 'Low', 'status `invalid` (SODOGI updated by new data)'],
        ['TRANS_ZMB_US', 'M-ZMB01', 'Medium', 'there has been a change in quantity from 0 to non-zero. this data will be processed once versioning is implemented'],
        ['TRANS_ZMB_US', 'M-ZMB02', 'Medium', 'affected by date filtering, so it will be processed during the monthly traceability'],
        ['TRANS_ZMB_US', 'M-ZMB03', 'Medium', 'ZMBUS that will be processed tomorrow or in the next few days'],
        ['TRANS_ZMB_US', 'L-ZMB01', 'Low', 'total = 0 (ZMB qty is 0)'],
        ['TRANS_ZMB_US', 'L-ZMB02', 'Low', 'process_information `Kirim...` or `TBS Kirim...` or `Stock...` or `No Mapping Yet`'],
        ['TRANS_OPENLINK_ESM', 'M-OPNK_ESM01', 'Medium', 'sap_so_po_number does not exist in sales_order_no trans_sodogi_esm'],
        ['TRANS_OPENLINK_ESM', 'M-OPNK_ESM02', 'Medium', 'sales_order_no trans_sodogi_esm connected to sap_so_po_number has no do_number'],
        ['TRANS_OPENLINK_ESM', 'M-OPNK_ESM03', 'Medium', 'sales_order_no trans_sodogi_esm connected to sap_so_po_number--has no do_number connected to reference trans_mb51_esm'],
        ['TRANS_OPENLINK_ESM', 'M-OPNK_ESM04', 'Medium', 'without bl_date'],
        ['TRANS_OPENLINK_ESM', 'M-OPNK_ESM05', 'Medium', 'without bl_number'],
        ['TRANS_OPENLINK_ESM', 'L-OPNK_ESM01', 'Low', 'Openlink ESM with sap_so_po_number startswith `35..` or `36..` is not used'],
        ['TRANS_SODOGI_ESM', 'M-SODOGI_ESM01', 'Medium', 'without do_number'],
        ['TRANS_SODOGI_ESM', 'M-SODOGI_ESM02', 'Medium', 'issued_qty (do_qty) = 0'],
        ['TRANS_SODOGI_ESM', 'M-SODOGI_ESM03', 'Medium', 'do_number is not connected to reference trans_mb51_esm'],
        ['TRANS_SODOGI_ESM', 'M-SODOGI_ESM04', 'Medium', 'without good_issue_date'],
        ['TRANS_MB51_ESM', 'M-MB51_ESM01', 'Medium', 'noise MB51 ESM'],
        ['TRANS_MB51_ESM', 'M-MB51_ESM02', 'Medium', 'not in main movement type MB51 ESM'],
        ['TRANS_MB51_ESM', 'M-MB51_ESM03', 'Medium', 'MB51ESM that will be process tomorrow'],
        ['TRANS_MB51_ESM', 'L-MB51_ESM01', 'Low', 'block-unblock process MB51 ESM'],
        ['-', 'H-UNKNOWN', 'High', 'unidentified'],
        ['-', 'H-RESTAN01', 'High', 'already included in restan and the quantity has been fully used, but the transaction has not yet been recorded in the events'],
        ['-', 'H-RESTAN02', 'High', 'already included in restan and the quantity has been partially used, but the transaction has not yet been recorded in the events'],
        ['-', 'M-RESTAN01', 'Medium', 'already included in restan, but there are no transactions linked to the restan data yet'],
        ['-', 'M-RESTAN02', 'Medium', 'already included in restan, but the restan data is not being used']
    ]
    df_internal_code = pd.DataFrame(internal_code, columns=['Task Name', 'Issue Code', 'Severity', 'Description'])
    return df_internal_code

#function to create glosary
def glosary():
    glosary = [
        ['noise', 'MB51DS/MB51US/MB51ESM data that is not used due to the cancellation of transaction data'],
        ['loss, gain, etc.', 'MB51DS/MB51US data that is not used because it is neither classified as noise data nor part of the odd movementtype'],
        ['block-unblock process', 'MB51DS/MB51ESM data with movementtype 343-344'],
        ['main movementtype MB51 ESM', '131;531;101;301;311;917;309;261;601;321;344;343,453;451;651;981;971;641;645'],
        ['unused material MB51US', '80000005 (PALM KERNEL SHELL); 80000006 (POME Oil); 80000007 (EFB Oil)'],
        ['unused movementtype MB51US for estate', 'all movementtype except 351; 961; 601'],
        ['plant mill ZMB unused', 'MB51US mill data with that is not used because movementtype not in 351, 961, 311, 601; and qtyinunofentry>0; and sloc is not contains CPO/PK']
    ]
    df_glosary = pd.DataFrame(glosary, columns=['Terms', 'Description'])
    return df_glosary

#function to send email with SES
def send_email_with_attachment(s3_client, s3_bucket, s3_key, recipients, date, log_etl, total_clean, total_disconnected, report_custom_date):

    perc_record_s3_raw = round(len(log_etl[log_etl['Status Record S3 to Raw']=='PASS']) * 100.0 / len(log_etl), 2)
    perc_column_s3_raw = round(len(log_etl[log_etl['Status Column S3 to Raw']=='PASS']) * 100.0 / len(log_etl), 2)
    perc_record_raw_clean = round(len(log_etl[log_etl['Status Record Raw to Clean']=='PASS']) * 100.0 / len(log_etl), 2)

    perc_traceability = (total_clean - total_disconnected) * 100 / total_clean
    perc_traceability = math.floor(perc_traceability * 100) / 100

    ses_client = boto3.client('ses', region_name='ap-southeast-3')
    sender = "cs@hara.ag"
    if report_custom_date == '':
        subject = f"ETL & Traceability Report, {date}"
    else:
        subject = f"[MANUAL] ETL & Traceability Report, {date}"
    body_text = f"""[Daily Summary ETL & Traceability Report]
    - Status Record S3 to Raw: {perc_record_s3_raw}% Pass
    - Status Column S3 to Raw: {perc_column_s3_raw}% Pass
    - Status Record Raw to Clean: {perc_record_raw_clean}% Pass
    - Status Transaction: {perc_traceability}% Processed"""

    try:
        report_content = download_file_from_s3(s3_client, s3_bucket, s3_key)

        msg = MIMEMultipart()
        msg['Subject'] = subject
        msg['From'] = sender
        msg['To'] = ", ".join(recipients)

        msg.attach(MIMEText(body_text, 'plain'))

        part = MIMEApplication(report_content)
        part.add_header('Content-Disposition', 'attachment', filename=os.path.basename(s3_key))
        msg.attach(part)

        response = ses_client.send_raw_email(
            Source=sender,
            Destinations=recipients,
            RawMessage={
                'Data': msg.as_string()
            }
        )
        print(f"Email sent! Message ID: {response['MessageId']}")

    except ClientError as e:
        print(f"Failed to send email: {e.response['Error']['Message']}")
        raise e

#function to send email with SES
def send_email_no_activity(recipients, date, report_custom_date):
    ses_client = boto3.client('ses', region_name='ap-southeast-3')

    sender = "cs@hara.ag"
    if report_custom_date == '':
        subject = f"ETL & Traceability Report, {date}"
    else:
        subject = f"[MANUAL] ETL & Traceability Report, {date}"
    body_text = f"There is no ETL activity on {date}."

    try:
        response = ses_client.send_email(
                    Source=sender,
                    Destination={
                        'ToAddresses': recipients  
                    },
                    Message={
                        'Subject': {
                            'Data': subject
                        },
                        'Body': {
                            'Text': {
                                'Data': body_text
                            }
                        }
                    })
        print(f"Email sent! Message ID: {response['MessageId']}")
    except ClientError as e:
        print(f"Failed to send email: {e.response['Error']['Message']}")
        raise e
